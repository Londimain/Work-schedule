# tab_report.py
# -*- coding: utf-8 -*-
"""Вкладка «Отчёт и печать» и экспорт в Excel."""
import config  # noqa: F401

import os
import webbrowser
from calendar import monthrange
from datetime import datetime, date
from ui_helpers import NoOverscrollEffect   # <— новое

from kivy.metrics import dp
from kivy.uix.anchorlayout import AnchorLayout
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.button import Button
from kivy.uix.label import Label
from kivy.uix.popup import Popup
from kivy.uix.scrollview import ScrollView
from kivy.uix.spinner import Spinner
from kivy.uix.widget import Widget
from kivy.graphics import Color, Rectangle
from kivy.clock import Clock

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
    OPENPYXL_ERR = ""
except Exception as _e:
    HAS_OPENPYXL = False
    OPENPYXL_ERR = f"{type(_e).__name__}: {_e}"
    import traceback
    OPENPYXL_ERR += "\n\n" + traceback.format_exc()

from config import (FONT_NAME, LABEL_LIGHT, MONTHS, YEARS,
                    SPINNER_BG, SPINNER_FG, RIGHT_COLUMNS,
                    WEEKDAYS_SHORT, DEFAULT_ORG, BASE_DIR,
                    DEFAULT_ENERGETIC, DEFAULT_DIRECTOR)
from data_utils import (save_data, classify_day, analyze_row,
                        apply_overtime, fmt_hours)
from ui_widgets import CenteredField


class ReportTab(BoxLayout):
    def __init__(self, app, **kwargs):
        super().__init__(orientation='vertical', **kwargs)
        self.app = app
        self._build()
        Clock.schedule_once(lambda *a: self.update_preview(), 0.3)

    def _build(self):
        top = BoxLayout(size_hint_y=None, height=dp(52), spacing=5, padding=5)

        top.add_widget(Label(text="Месяц:", size_hint_x=None, width=dp(70),
                             color=LABEL_LIGHT, font_name=FONT_NAME,
                             font_size=dp(15)))
        self.month_spinner = Spinner(text=MONTHS[datetime.now().month - 1],
                                     values=MONTHS, size_hint_x=None, width=dp(150),
                                     font_name=FONT_NAME,
                                     background_color=SPINNER_BG, color=SPINNER_FG)
        top.add_widget(self.month_spinner)

        top.add_widget(Label(text="Год:", size_hint_x=None, width=dp(50),
                             color=LABEL_LIGHT, font_name=FONT_NAME,
                             font_size=dp(15)))
        self.year_spinner = Spinner(text=str(datetime.now().year),
                                    values=YEARS, size_hint_x=None, width=dp(100),
                                    font_name=FONT_NAME,
                                    background_color=SPINNER_BG, color=SPINNER_FG)
        top.add_widget(self.year_spinner)

        top.add_widget(Widget())  # распорка

        excel_btn = Button(text="Сохранить в Excel",
                           size_hint_x=None, width=dp(200),
                           font_name=FONT_NAME, font_size=dp(15), bold=True,
                           background_color=(0.25, 0.55, 0.30, 1),
                           color=(1, 1, 1, 1))
        excel_btn.bind(on_press=self.save_excel)
        top.add_widget(excel_btn)
        self.add_widget(top)

        self.month_spinner.bind(text=lambda *a: self.update_preview())
        self.year_spinner.bind(text=lambda *a: self.update_preview())

        line = Widget(size_hint_y=None, height=dp(2))
        with line.canvas:
            Color(0.5, 0.5, 0.5, 1)
            line._rect = Rectangle(pos=line.pos, size=line.size)
        line.bind(pos=lambda i, v: setattr(i._rect, 'pos', v),
                  size=lambda i, v: setattr(i._rect, 'size', v))
        self.add_widget(line)

        scroll = ScrollView(scroll_type=['bars'],
                            effect_cls=NoOverscrollEffect)
        self.preview = Label(text="", size_hint_y=None, halign='left',
                              valign='top', padding=(20, 20),
                              font_name=FONT_NAME,
                              font_size=dp(13), color=LABEL_LIGHT)
        self.preview.bind(width=lambda *a: setattr(
            self.preview, 'text_size', (self.preview.width - 40, None)))
        self.preview.bind(texture_size=self.preview.setter('size'))
        scroll.add_widget(self.preview)
        self.add_widget(scroll)

    # ---------------- энергетик -------------------------------
    def _get_signer(self, dept):
        """Возвращает (должность, ФИО) подписанта для отдела."""
        signers = self.app.data.get('department_signer', {}) or {}
        entry = signers.get(dept)
        if entry and entry.get('name'):
            return (entry.get('position') or 'Энергетик',
                    entry.get('name'))
        # legacy-fallback (старые данные)
        em = self.app.data.get('department_energetic', {}) or {}
        name = em.get(dept) or self.app.data.get('energetic_name',
                                                 DEFAULT_ENERGETIC)
        return ('Энергетик', name)

    # ---------------- вспомогательные ------------------------------------
    def _employees_for_report(self):
        return [emp for emp in self.app.data['employees']
                if emp.get('show_in_timesheet', False)]

    def _get_dept(self):
        emps = self._employees_for_report()
        if emps:
            return emps[0].get('department', '')
        return ''

    def _calc_grand(self):
        y = int(self.year_spinner.text)
        m = MONTHS.index(self.month_spinner.text) + 1
        key = f"{y}-{m:02d}"
        days_in_month = monthrange(y, m)[1]
        ts = self.app.data['timesheets'].get(key, {})
        employees = self._employees_for_report()
        night_flags = self.app.data.get('night_flags', {})
        overtime_flags = self.app.data.get('overtime_flags', {})
        holiday_flags = self.app.data.get('holiday_flags', {})
        day_types = [classify_day(y, m, d)
                     for d in range(1, days_in_month + 1)]
        work_days = sum(1 for dt in day_types if dt == 'work')
        norm_hours = work_days * 8.0
        overrides_all = self.app.data.get('result_overrides', {}).get(key, {})

        grand = {k: 0 for _, k, _ in RIGHT_COLUMNS}
        grand['часы'] = 0.0
        grand['сверхурочные'] = 0.0
        grand['ночные'] = 0.0
        grand['праздничные'] = 0.0

        for emp in employees:
            emp_id = emp['id']
            day_data = ts.get(emp_id, {})
            values = [day_data.get(str(d), '')
                      for d in range(1, days_in_month + 1)]
            pos_name = emp.get('position', '')
            night_on = bool(night_flags.get(pos_name, False))
            overtime_on = bool(overtime_flags.get(pos_name, False))
            holiday_on = bool(holiday_flags.get(pos_name, False))
            res = analyze_row(values, day_types, night_on, holiday_on)
            if overtime_on:
                apply_overtime(res, norm_hours)
            else:
                res['сверхурочные'] = 0.0
            emp_over = overrides_all.get(emp_id, {})
            for k in grand:
                if k in emp_over:
                    try:
                        grand[k] += float(str(emp_over[k]).replace(',', '.'))
                    except (ValueError, TypeError):
                        pass
                else:
                    grand[k] += res.get(k, 0)
        return grand, days_in_month, work_days, norm_hours, key

    # ---------------- предпросмотр ---------------------------------------
    def update_preview(self):
        y = int(self.year_spinner.text)
        m = MONTHS.index(self.month_spinner.text) + 1
        key = f"{y}-{m:02d}"
        days_in_month = monthrange(y, m)[1]
        ts = self.app.data['timesheets'].get(key, {})
        employees = self._employees_for_report()
        night_flags = self.app.data.get('night_flags', {})
        overtime_flags = self.app.data.get('overtime_flags', {})
        holiday_flags = self.app.data.get('holiday_flags', {})
        org = self.app.data.get('organization', DEFAULT_ORG)
        dept = self._get_dept()
        signer_pos, signer_name = self._get_signer(dept)
        director_name = self.app.data.get('director_name', DEFAULT_DIRECTOR)

        if not employees:
            self.preview.text = "Нет сотрудников, отмеченных «В табель»."
            return

        grand, days_in_month, work_days, norm_hours, key = self._calc_grand()

        sep = "─" * 90
        lines = [
            "  ТАБЕЛЬ УЧЁТА РАБОЧЕГО ВРЕМЕНИ",
            f"  Период: {MONTHS[m-1]} {y}   ({days_in_month} календарных дней)",
            f"  Организация: {org}",
            f"  Отдел: {dept}",
            f"  Норма: {work_days} будних × 8 = {fmt_hours(norm_hours)} ч",
            sep,
            f"  Сотрудников: {len(employees)}",
            f"  Факт. работы: {grand.get('явки', 0)}    "
            f"Выходных: {grand.get('вых', 0)}    "
            f"Труд. отпуск: {grand.get('отп', 0)}    "
            f"Болезнь: {grand.get('бол', 0)}    "
            f"За свой счёт: {grand.get('за_свой', 0)}",
            f"  Прогул: {grand.get('прогул', 0)}    "
            f"Текущий простой: {grand.get('текущий_простой', 0)}    "
            f"Опоздания: {grand.get('опозд', 0)}    "
            f"Целосмен.простой: {grand.get('целосмен_простой', 0)}",
            f"  Роды: {grand.get('отп_роды', 0)}    "
            f"Проч.неявки: {grand.get('проч_неявки', 0)}",
            f"  Всего часов: {fmt_hours(grand.get('часы', 0))}    "
            f"Сверхурочные: {fmt_hours(grand.get('сверхурочные', 0))}    "
            f"Ночные: {fmt_hours(grand.get('ночные', 0))}    "
            f"Праздничные: {fmt_hours(grand.get('праздничные', 0))}",
            f"  Подписи: {signer_pos} — {signer_name}    "
            f"Директор — {director_name}",
            sep,
            "",
            "  По сотрудникам:",
        ]

        name_w = 26
        day_types = [classify_day(y, m, d)
                     for d in range(1, days_in_month + 1)]
        for idx, emp in enumerate(employees):
            emp_id = emp['id']
            day_data = ts.get(emp_id, {})
            values = [day_data.get(str(d), '')
                      for d in range(1, days_in_month + 1)]
            pos_name = emp.get('position', '')
            night_on = bool(night_flags.get(pos_name, False))
            overtime_on = bool(overtime_flags.get(pos_name, False))
            holiday_on = bool(holiday_flags.get(pos_name, False))
            res = analyze_row(values, day_types, night_on, holiday_on)
            if overtime_on:
                apply_overtime(res, norm_hours)
            else:
                res['сверхурочные'] = 0.0
            nm = (emp.get('name', '') or '')[:name_w].ljust(name_w)
            lines.append(
                f"   {idx + 1:>2}. {nm} "
                f"явок {res['явки']:>2}  часов {fmt_hours(res['часы']):>7}  "
                f"сверхур {fmt_hours(res['сверхурочные']):>5}  "
                f"празд {fmt_hours(res['праздничные']):>5}  "
                f"В {res['вых']:>2}  О {res['отп']:>2}  Б {res['бол']:>2}"
            )

        lines += ["", sep,
                  "  Нажмите «Сохранить в Excel», чтобы получить файл .xlsx"]
        self.preview.text = "\n".join(lines)

    # ---------------- экспорт в Excel ------------------------------------
    def save_excel(self, *a):
        if not HAS_OPENPYXL:
            from kivy.uix.scrollview import ScrollView as _SV
            msg = ("Не удалось загрузить openpyxl в собранной программе.\n\n"
                   "Техническая информация:\n"
                   f"{OPENPYXL_ERR}")
            inner = BoxLayout(orientation='vertical', spacing=8, padding=8)
            sv = _SV()
            lbl = Label(text=msg, color=LABEL_LIGHT, halign='left',
                        valign='top', font_name=FONT_NAME, font_size=dp(13))
            lbl.bind(width=lambda i, v: setattr(i, 'text_size', (v, None)))
            lbl.bind(texture_size=lbl.setter('size'))
            sv.add_widget(lbl)
            inner.add_widget(sv)
            ok_btn = Button(text="OK", size_hint_y=None, height=dp(40),
                            font_name=FONT_NAME, font_size=dp(16), bold=True,
                            background_color=(0.3, 0.7, 0.3, 1),
                            color=(1, 1, 1, 1))
            inner.add_widget(ok_btn)
            popup = Popup(title="Нужен openpyxl", content=inner,
                          size_hint=(0.85, 0.55),
                          title_color=LABEL_LIGHT, title_font=FONT_NAME)
            ok_btn.bind(on_press=popup.dismiss)
            popup.open()
            return

        y = int(self.year_spinner.text)
        m = MONTHS.index(self.month_spinner.text) + 1
        key = f"{y}-{m:02d}"
        days_in_month = monthrange(y, m)[1]
        ts = self.app.data['timesheets'].get(key, {})
        employees = self._employees_for_report()
        night_flags = self.app.data.get('night_flags', {})
        overtime_flags = self.app.data.get('overtime_flags', {})
        holiday_flags = self.app.data.get('holiday_flags', {})
        org = self.app.data.get('organization', DEFAULT_ORG)
        dept = self._get_dept()
        signer_pos, signer_name = self._get_signer(dept)
        director_name = self.app.data.get('director_name', DEFAULT_DIRECTOR)

        day_types = [classify_day(y, m, d)
                     for d in range(1, days_in_month + 1)]
        work_days = sum(1 for dt in day_types if dt == 'work')
        norm_hours = work_days * 8.0
        overrides_all = self.app.data.get('result_overrides', {}).get(key, {})

        thin = Side(border_style="thin", color="333333")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal='center', vertical='center')
        left = Alignment(horizontal='left', vertical='center')
        vert = Alignment(horizontal='center', vertical='center',
                         text_rotation=90, wrap_text=False)
        header_fill = PatternFill("solid", fgColor="E8E8E8")
        weekend_fill = PatternFill("solid", fgColor="FFDEDE")
        holiday_fill = PatternFill("solid", fgColor="FFF3CC")
        total_fill = PatternFill("solid", fgColor="D0D0D0")
        reg9 = Font(size=9)
        reg14 = Font(size=14)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{MONTHS[m-1]} {y}"

        total_cols = 3 + days_in_month + len(RIGHT_COLUMNS)

        ws.merge_cells(start_row=1, start_column=1,
                       end_row=1, end_column=total_cols)
        c = ws.cell(row=1, column=1, value="ТАБЕЛЬ УЧЁТА РАБОЧЕГО ВРЕМЕНИ")
        c.font = reg14
        c.alignment = center

        ws.merge_cells(start_row=2, start_column=1,
                       end_row=2, end_column=total_cols)
        c = ws.cell(row=2, column=1,
                    value=(f'Организация: {org} | Отдел: {dept} | '
                           f'Период: {MONTHS[m-1]} {y} '
                           f'({days_in_month} календарных дней)'))
        c.alignment = center
        c.font = Font(size=10)

        row_h1, row_h2 = 3, 4
        for col, label in [(1, "№"), (2, "Фамилия Имя Отчество"),
                           (3, "Должность")]:
            c = ws.cell(row=row_h1, column=col, value=label)
            c.font = reg9
            c.alignment = center
            c.fill = header_fill
            ws.merge_cells(start_row=row_h1, start_column=col,
                           end_row=row_h2, end_column=col)

        start_days = 4
        end_days = start_days + days_in_month - 1
        ws.merge_cells(start_row=row_h1, start_column=start_days,
                       end_row=row_h1, end_column=end_days)
        c = ws.cell(row=row_h1, column=start_days, value="Числа месяца")
        c.font = reg9
        c.alignment = center
        c.fill = header_fill

        start_r = end_days + 1
        for span, title in [(2, "Дни явок"),
                            (9, "Дни неявок"),
                            (4, "Отработано часов")]:
            ws.merge_cells(start_row=row_h1, start_column=start_r,
                           end_row=row_h1, end_column=start_r + span - 1)
            c = ws.cell(row=row_h1, column=start_r, value=title)
            c.font = reg9
            c.alignment = center
            c.fill = header_fill
            start_r += span

        for i, dt in enumerate(day_types):
            col = start_days + i
            wd = WEEKDAYS_SHORT[date(y, m, i + 1).weekday()]
            c = ws.cell(row=row_h2, column=col, value=f"{i + 1}\n{wd}")
            c.font = Font(size=9)
            c.alignment = Alignment(horizontal='center', vertical='top',
                                    wrap_text=True)
            if dt == 'weekend':
                c.fill = weekend_fill
            elif dt == 'holiday':
                c.fill = holiday_fill
            else:
                c.fill = header_fill

        start_r = end_days + 1
        for label, key_, _ in RIGHT_COLUMNS:
            c = ws.cell(row=row_h2, column=start_r, value=label)
            c.font = Font(size=10)
            c.alignment = vert
            c.fill = header_fill
            start_r += 1

        for r in (row_h1, row_h2):
            for col in range(1, total_cols + 1):
                ws.cell(row=r, column=col).border = border

        row = row_h2 + 1
        grand = {k: 0 for _, k, _ in RIGHT_COLUMNS}
        grand['часы'] = 0.0
        grand['сверхурочные'] = 0.0
        grand['ночные'] = 0.0
        grand['праздничные'] = 0.0

        for idx, emp in enumerate(employees):
            emp_id = emp['id']
            day_data = ts.get(emp_id, {})
            values = [day_data.get(str(d), '')
                      for d in range(1, days_in_month + 1)]
            pos_name = emp.get('position', '')
            night_on = bool(night_flags.get(pos_name, False))
            overtime_on = bool(overtime_flags.get(pos_name, False))
            holiday_on = bool(holiday_flags.get(pos_name, False))
            res = analyze_row(values, day_types, night_on, holiday_on)
            if overtime_on:
                apply_overtime(res, norm_hours)
            else:
                res['сверхурочные'] = 0.0
            emp_over = overrides_all.get(emp_id, {})

            for k in grand:
                if k in emp_over:
                    try:
                        grand[k] += float(str(emp_over[k]).replace(',', '.'))
                    except (ValueError, TypeError):
                        pass
                else:
                    grand[k] += res.get(k, 0)

            ws.cell(row=row, column=1, value=idx + 1).alignment = center
            ws.cell(row=row, column=2, value=emp.get('name', '')).alignment = left
            ws.cell(row=row, column=3, value=emp.get('position', '')).alignment = left

            for i, v in enumerate(values):
                col = start_days + i
                c = ws.cell(row=row, column=col, value=v)
                c.alignment = center
                if day_types[i] == 'weekend':
                    c.fill = weekend_fill
                elif day_types[i] == 'holiday':
                    c.fill = holiday_fill

            start_r = end_days + 1
            for _, key_, _ in RIGHT_COLUMNS:
                if key_ in emp_over:
                    val = emp_over[key_]
                else:
                    v = res.get(key_, 0)
                    val = (fmt_hours(v) if key_ in
                           ('часы', 'сверхурочные', 'ночные', 'праздничные')
                           else v)
                c = ws.cell(row=row, column=start_r, value=val)
                c.alignment = center
                start_r += 1

            for col in range(1, total_cols + 1):
                ws.cell(row=row, column=col).border = border
            row += 1

        ws.merge_cells(start_row=row, start_column=1,
                       end_row=row, end_column=3)
        ws.cell(row=row, column=1, value="ИТОГО").font = Font()
        ws.cell(row=row, column=1).alignment = center
        for col in range(1, 4):
            ws.cell(row=row, column=col).border = border
            ws.cell(row=row, column=col).fill = total_fill

        ws.merge_cells(start_row=row, start_column=start_days,
                       end_row=row, end_column=end_days)
        for col in range(start_days, end_days + 1):
            ws.cell(row=row, column=col).border = border
            ws.cell(row=row, column=col).fill = total_fill

        start_r = end_days + 1
        for _, key_, _ in RIGHT_COLUMNS:
            v = grand.get(key_, 0)
            val = (fmt_hours(v) if key_ in
                   ('часы', 'сверхурочные', 'ночные', 'праздничные')
                   else v)
            c = ws.cell(row=row, column=start_r, value=val)
            c.font = Font()
            c.alignment = center
            c.border = border
            c.fill = total_fill
            start_r += 1

        max_name_len = max((len(emp.get('name', '') or '')
                            for emp in employees), default=20)
        max_pos_len  = max((len(emp.get('position', '') or '')
                            for emp in employees), default=15)

        ws.column_dimensions['A'].width = 5
        ws.column_dimensions['B'].width = max(28, max_name_len + 2)
        ws.column_dimensions['C'].width = max(20, max_pos_len + 2)
        for i in range(days_in_month):
            ws.column_dimensions[get_column_letter(start_days + i)].width = 3.5
        for i in range(len(RIGHT_COLUMNS)):
            ws.column_dimensions[get_column_letter(end_days + 1 + i)].width = 6

        ws.row_dimensions[1].height = 22
        ws.row_dimensions[2].height = 16
        ws.row_dimensions[row_h2].height = 95

        row += 5          # отступ от таблицы до блока подписей

        # Слева — Руководитель, справа — Утверждение.
        # LEFT_COL — сразу у левого края таблицы (колонка C).
        # RIGHT_COL — ближе к правому краю (последние ~10 колонок свободны).
        LEFT_COL  = 3
        RIGHT_COL = max(LEFT_COL + 22, total_cols - 11)

        # --- Строка 1: метки ---
        c = ws.cell(row=row, column=LEFT_COL, value="Руководитель:")
        c.font = Font(size=11)
        c.alignment = Alignment(horizontal='left', vertical='center')

        c = ws.cell(row=row, column=RIGHT_COL, value="Утверждение:")
        c.font = Font(size=11)
        c.alignment = Alignment(horizontal='left', vertical='center')

        # --- Строка 2: сами подписи (через 2 строки ниже) ---
        sig_row = row + 2

        c = ws.cell(row=sig_row, column=LEFT_COL,
                    value=f"{signer_pos}   __________________   {signer_name}")
        c.font = Font(size=11)
        c.alignment = Alignment(horizontal='left', vertical='center',
                                 indent=12)

        c = ws.cell(row=sig_row, column=RIGHT_COL,
                    value=f"Директор   __________________   {director_name}")
        c.font = Font(size=11)
        c.alignment = Alignment(horizontal='left', vertical='center',
                                 indent=12)

        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        # Область печати — до нижней строки подписей + 2 строки запаса
        last_print_row = sig_row + 2
        try:
            ws.print_area = (f"A1:"
                             f"{get_column_letter(total_cols)}"
                             f"{last_print_row}")
        except Exception:
            pass

        def _sanitize_folder(name):
            bad = '\\/:*?"<>|'
            for ch in bad:
                name = name.replace(ch, '_')
            name = name.strip()
            return name or 'Без отдела'

        dept_folder = _sanitize_folder(dept)
        out_dir = os.path.join(BASE_DIR, "Табель", dept_folder, str(y))
        try:
            os.makedirs(out_dir, exist_ok=True)
        except Exception as e:
            print("Не удалось создать папку:", e)

        path = os.path.join(out_dir, f"tabel_{key}.xlsx")

        try:
            wb.save(path)
        except Exception as e:
            print("Ошибка сохранения Excel:", e)
            return
        try:
            os.startfile(path)
        except Exception:
            try:
                webbrowser.open(f"file://{path}")
            except Exception:
                pass